"""Render the vector/text research deck to PDF using headless Chrome."""

from __future__ import annotations

import html
import shutil
import tempfile
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "endstep_tournament_report.pptx"
PDF = ROOT / "endstep_tournament_report.pdf"
EMU_PER_INCH = 914400


def inches(value: int) -> float:
    return value / EMU_PER_INCH


def color(value, fallback: str = "transparent") -> str:
    try:
        rgb = value.rgb
        return f"#{rgb}"
    except (AttributeError, TypeError):
        return fallback


def shape_fill(shape) -> str:
    try:
        return color(shape.fill.fore_color, "transparent")
    except (AttributeError, TypeError):
        return "transparent"


def shape_line(shape) -> tuple[str, float]:
    try:
        return color(shape.line.color, "transparent"), shape.line.width / 12700
    except (AttributeError, TypeError, ValueError):
        return "transparent", 0


def text_style(shape) -> tuple[str, float, bool, str]:
    font_name = "Aptos"
    font_size = 12.0
    bold = False
    text_color = "#F7F3E8"
    if not getattr(shape, "has_text_frame", False):
        return font_name, font_size, bold, text_color
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.name:
                font_name = run.font.name
            if run.font.size:
                font_size = run.font.size.pt
            if run.font.bold is not None:
                bold = run.font.bold
            try:
                text_color = color(run.font.color, text_color)
            except AttributeError:
                pass
            return font_name, font_size, bold, text_color
    return font_name, font_size, bold, text_color


def alignment(shape) -> str:
    try:
        value = shape.text_frame.paragraphs[0].alignment
    except (AttributeError, IndexError):
        return "left"
    if value == PP_ALIGN.CENTER:
        return "center"
    if value == PP_ALIGN.RIGHT:
        return "right"
    return "left"


def render_shape(shape) -> str:
    left, top = inches(shape.left), inches(shape.top)
    width, height = inches(shape.width), inches(shape.height)
    fill = shape_fill(shape)
    line, line_width = shape_line(shape)
    radius = "0.16in" if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE else "0"
    style = (
        f"left:{left:.4f}in;top:{top:.4f}in;width:{width:.4f}in;height:{height:.4f}in;"
        f"background:{fill};border:{line_width:.2f}pt solid {line};border-radius:{radius};"
    )
    if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
        return f'<div class="shape" style="{style}"></div>'

    font_name, font_size, bold, text_color = text_style(shape)
    content = html.escape(shape.text).replace("\n", "<br>")
    text_css = (
        f"font-family:{html.escape(font_name)};font-size:{font_size:.2f}pt;"
        f"font-weight:{700 if bold else 400};color:{text_color};text-align:{alignment(shape)};"
    )
    return f'<div class="shape text" style="{style}{text_css}">{content}</div>'


def build_html() -> str:
    prs = Presentation(PPTX)
    slide_width = inches(prs.slide_width)
    slide_height = inches(prs.slide_height)
    slides = []
    for slide in prs.slides:
        background = "#10141C"
        try:
            background = color(slide.background.fill.fore_color, background)
        except (AttributeError, TypeError):
            pass
        shapes = "\n".join(render_shape(shape) for shape in slide.shapes)
        slides.append(f'<section class="slide" style="background:{background};">{shapes}</section>')
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><style>
@page {{ size: {slide_width:.4f}in {slide_height:.4f}in; margin: 0; }}
* {{ box-sizing: border-box; }}
html, body {{ margin: 0; padding: 0; background: #10141C; }}
.slide {{ position: relative; display: block; width: {slide_width:.4f}in; height: {slide_height:.4f}in;
  overflow: hidden; break-after: page; page-break-after: always; page-break-inside: avoid; print-color-adjust: exact; }}
.slide + .slide {{ break-before: page; page-break-before: always; }}
.slide:last-child {{ break-after: auto; page-break-after: auto; }}
.shape {{ position: absolute; overflow: hidden; }}
.text {{ padding: 0.02in 0.03in; line-height: 1.08; white-space: normal; }}
</style></head><body>{"".join(slides)}</body></html>"""


def chrome_path() -> str:
    candidates = [
        shutil.which("google-chrome"),
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
        "/Applications/Chromium.app/Contents/MacOS/Chromium",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return str(candidate)
    raise RuntimeError("Google Chrome or Chromium is required to render the PDF")


def build() -> None:
    with tempfile.TemporaryDirectory(prefix="endstep-pdf-") as temp_dir:
        html_path = Path(temp_dir) / "deck.html"
        html_path.write_text(build_html(), encoding="utf-8")
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(executable_path=chrome_path(), headless=True)
            page = browser.new_page()
            page.goto(html_path.as_uri(), wait_until="load")
            page.pdf(
                path=str(PDF),
                print_background=True,
                prefer_css_page_size=True,
                display_header_footer=False,
            )
            browser.close()
    print(PDF)


if __name__ == "__main__":
    build()
